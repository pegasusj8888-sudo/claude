# research/hyundai/rows.jsonl → 현대_트랜스폰더_칩코드_DB.xlsx (기아 트랜스폰더 DB 구조 + 키종류, 맨 앞 브랜드)
import json, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(ROOT, '현대_트랜스폰더_칩코드_DB.xlsx')
rows = [json.loads(l) for l in open(os.path.join(HERE, 'rows.jsonl'), encoding='utf-8')]


def xt(chip):
    chip = chip or ''
    if not chip or (chip.startswith(('확인', '해당')) and 'ID' not in chip): return ('', '', '')
    new = any(k in chip for k in ('ID4A', 'ID47', 'ID49', 'ID8A', 'ID6A', 'ID88', 'AES'))
    old = any(k in chip for k in ('ID46', 'ID44', 'ID60', 'ID70', 'DST80', '4D70', '4D60x80', 'ID6E', 'ID4C'))
    if 'ID48' in chip:
        return ('△', '△', 'O') if old else ('△', '△', '△')
    if old and new: return ('△', '△', 'O')
    if old: return ('O', 'O', 'O')
    if new: return ('X', '△', 'O')
    return ('', '', '')


HDR = ['브랜드', '모델명', '연식', 'XT27A/A66 호환', 'XT27B 호환', 'XT57B 호환', '칩코드 (예: ID46(PCF7936))', '키종류',
       '키블레이드(부품번호)', '키블레이드(키웨이)', '카드키(부품번호)', '스마트키(부품번호)', '폴딩키(부품번호)',
       '이모빌라이저 시스템', '출처', '비고']
wb = openpyxl.Workbook(); ws = wb.active; ws.title = '현대 트랜스폰더 DB'
thin = Side(style='thin', color='FFBFBFBF'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ws.append(HDR)
for c in ws[1]:
    c.font = Font(name='Arial', size=11, bold=True, color='FFFFFFFF'); c.fill = PatternFill('solid', fgColor='FF002C5F')
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False); c.border = bd
F = lambda c: PatternFill('solid', fgColor=c)
XT = {'O': (F('FFC6EFCE'), 'FF006100'), '△': (F('FFFFEB9C'), 'FF9C6500'), 'X': (F('FFFFC7CE'), 'FF9C0006')}
FLAG = {'orange': (F('FFFFD599'), 'FF000000'), 'red': (F('FFFFC7CE'), 'FF9C0006')}
for r in rows:
    a, b, c = xt(r['chip'])
    ws.append([r['brand'], r['model'], r['year'], a, b, c, r['chip'], r['ktype'], r['blade_pn'] or None, r['keyway'] or None,
               r['card'] or None, r['smart'] or None, r['fold'] or None, r['immo'] or None, r['src'], r['note'] or None])
    i = ws.max_row
    for cell in ws[i]:
        cell.font = Font(name='Arial', size=10); cell.border = bd; cell.alignment = Alignment(vertical='center', wrap_text=False)
    for col in (4, 5, 6):
        cell = ws.cell(i, col); cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
        if cell.value in XT:
            cell.fill, fc = XT[cell.value]; cell.font = Font(name='Arial', size=10, bold=True, color=fc)
    fl = r.get('flag')
    if fl:
        fill, fc = FLAG[fl]
        for col in (7, 14, 16):
            ws.cell(i, col).fill = fill; ws.cell(i, col).font = Font(name='Arial', size=10, color=fc)
for col, w in zip('ABCDEFGHIJKLMNOP', [10, 34, 12, 12, 11, 11, 40, 20, 40, 12, 30, 60, 40, 14, 60, 60]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'C2'; ws.auto_filter.ref = f'A1:P{ws.max_row}'

g = wb.create_sheet('안내')
notes = [
'이 표는 락스미스(자동차 키 제작/프로그래밍) 참고용입니다.', '',
'※ 구성',
"- 예전 현대 파일(컬럼 10개, 칩·세부ID·부품번호가 섞인 구조)을 버리고 '기아_트랜스폰더_칩코드_DBnew.xlsx'와 같은 컬럼으로 새로 작성했습니다. 칩코드 뒤에 '키종류'를 추가했습니다.",
"- 맨 앞 '브랜드' 컬럼에 브랜드명을 적었습니다(검색 프로그램에서 '현대 쏘나타', '제네시스 G80'처럼 브랜드와 모델명을 함께 검색 가능). 제네시스 브랜드 독립(2015.11) 이후 모델은 '제네시스', 그 이전 현대 제네시스(BH·DH)·제네시스 쿠페는 '현대'입니다.",
"- 모델마다 연식별로 한 행씩, 2000년식부터 기록했습니다. 예전 파일에서 연식을 나누지 않고 한 줄로 적었던 모델(예: 2006-2010)은 조사 대상에서 뺐습니다(에쿠스·i20·갤로퍼·테라칸·트라제 XG·라비타·클릭·마이티 등).",
"- 연식은 실제 국내 판매(출시일) 기준입니다. 풀체인지가 있는 해는 전·후 세대가 각각 한 행씩 있습니다. 후속 없이 단종된 모델의 마지막 연식은 '2019(단종)' 형식입니다.",
"- '키종류'는 그 연식에 나온 키 형태(막대키·폴딩키·스마트키·카드키)를 모두 적었습니다. 카드키는 그랜저 HG·IG, 제네시스 DH의 카드형 스마트키입니다.",
"- 키 부품번호는 국내 사양(433/434/447MHz)만 적었습니다. 미국형 315MHz 키(예: 95440-3Q000·95440-3M220·95440-2M350·95440-2V100·95440-4Z200, FCC SY5HMFNA04)는 뺐습니다.",
"- 스마트키에 딸린 비상키(블랭킹키)는 '키블레이드(부품번호)' 칸에 적었습니다. 키리모컨 일체형(81996-4H100 등)은 '폴딩키(부품번호)' 칸에 적었습니다.",
"- '이모빌라이저 시스템' 칸은 현대모비스 품번이 확인된 차종만 적었습니다(그랜드 스타렉스·포터2 이모빌라이저 모듈 95420-D4000·95420-4F500, 포터2 안테나 95440-4F100).",
"- 출처 칸은 검색에서 근거로 쓴 페이지입니다. hellowcar·partsro(현대모비스 공식 부품몰)는 적용 기간이 적힌 국내 순정 품번의 근거이고, 칩 종류는 대부분 해외 키 판매처(abkeys·mk3·remotesandkeys 등)의 같은 품번 설명을 근거로 했습니다. 원문은 이 환경에서 열 수 없어 검색 결과 요약을 근거로 했습니다.",
"- '비고'에는 주황색·빨간색 칸의 이유만 적었습니다.",
'',
'※ 색상',
"- 주황색: 추정이거나 자료가 서로 다른 값 — 실물 키·VIN으로 재확인하세요.",
"- 빨간색: 확인 불가 — 아반떼 J2(2000) 1행.",
"- 주황색으로 남은 행(재조사 후 23행): 부분변경 신형 스마트키 칩 원문이 아직 없는 경우(아이오닉6 2025~, 일렉트리파이드 G80 2024~, 일렉트리파이드 GV70 2025~, GV60 2026, 포터2 LPG 스마트키, 포터2 일렉트릭 CN000), 스타렉스 2004-2005(순정 4D60 키 vs 해외 카탈로그 4C), 베르나 LC(이모빌라이저키 품번이 없어 미적용 추정).",
"- 전환 연식은 칩코드 칸에 전·후를 함께 적었습니다(예: 쏘나타 DN8 2023 = ID47(NCF29A1X)[~2023.2], ID4A(NCF29A1M)[디 엣지, 2023.3~]).",
"- 칩코드 추가 뜻: ID4C = Texas 4C(고정 코드, 스타렉스 2004-2006 해외 자료), 해당 없음 = 그 연식 국내형에 이모빌라이저가 없음.",
'',
'※ 예전 현대 파일에서 바로잡은 주요 오류',
"- 쏘나타 LF(2014-2019): ID47 → ID8A(DST-AES). 95440-C1001·C1500 모두 8A.",
"- 쏘나타 DN8 2019-2022: ID4A → ID47(NCF29A1X). ID4A는 디 엣지(2023.3~)부터.",
"- 아반떼 AD(2015-2020): ID47 → ID8A. 아반떼 CN7·아반떼 N·아이오닉6: ID4A → ID6A. 더 뉴 아반떼(95440-AA500)·아이오닉6(KL000)은 FCC NYOMBEC7FOB2208(MBEC 계열 = CN7 AA000·캐스퍼와 같은 6A 계열)이고 auto-keys·keyshop-online·abkeys가 6A로 일치. 일부 판매처(remotesandkeys)의 '4A' 표기는 같은 키를 4A·6A로 섞어 적어 따르지 않았습니다.",
"- 투싼 NX4 부분변경(더 뉴 투싼, 2023.11~): 스마트키가 95440-N95xx(ID4A)로 바뀜 — 처음 작성 때 ID47로 잘못 이어 적었던 것을 정정.",
"- 그랜저 GN7: 스마트키 품번 95440-BY000 → 95440-N1xxx(N1000·N1050·N1110 등).",
"- 제네시스 DH 2013-2014: 'BH 전자장비 유지, ID46(95440-3M220)' → DH는 출시(2013.11)부터 95440-B1200BLH ID47. 95440-3M220은 미국형 315MHz.",
"- 제네시스 BH: 95440-3M220(미국형) → 국내 95440-3M010·3M020·3M030. 제네시스 쿠페: 95440-2M350·2M300·3V021(미국형) → 국내 95440-2M050·2M420.",
"- 싼타페 DM 2016-2018·맥스크루즈 2016-2019: 'ID46 또는 ID47' → ID46(95440-2W500·B8100 모두 PCF7952A). 싼타페 DM 95440-4Z200은 미국형.",
"- 싼타페 SM 2000-2002: 4D60 → 국내 이모빌라이저는 2005년형(2004.8.11)부터 적용, 그 이전은 해당 없음.",
"- 투싼 TL: 'ID6E-MA 또는 ID47' → ID47(NCF2952X).",
"- 스타리아: ID4A → ID47(NCF29A1X).",
"- 그랜드 스타렉스 2007-2011·포터2 2004-2015: '4D60 또는 미탑재' → 국내 이모빌라이저 사양 키는 ID46(그랜드 스타렉스 블랭킹키 81996-2H010 = 아반떼 HD 공용, 포터2 81996-4F020 = ID46 키블레이드).",
"- 베라크루즈 2013-2015: '수출형 ix55 확인 불가' → 국내 판매가 2015년까지 이어짐(95440-3J560 ID46).",
"- i40: 품번별로 칩이 다름 — 95440-3Z000(2011.12.15~2012.4.2)은 TMS37126(DST80, ID6E-MA), 3Z001·3Z002 이후는 DST-AES(8A).",
"- 그랜저 TG 스마트키 95440-3L100은 TIRIS DST80(ID6E-MA), 이모빌라이저키 81996-3L010은 ID46.",
"- 쏘나타 EF·그랜저 XG·싼타페 SM(2005년형~): 현대모비스 순정 이모빌라이저키 81996-38010(4D60). 베르나 MC: 81996-1E010(ID46, 트랜스폼 2009.6부터 이모빌라이저 삭제). 투싼 JM: 81996-2E010(ID46). 스타렉스: 81996-4A250(4D60).",
"- 포터2 스마트키는 2024 포터II LPG(2023.12~)부터(95440-4FGA0). 2020-2023 디젤에는 스마트키가 없어 예전 파일의 95440-4FGA0 기재는 오류.",
"- 코나 N 95440-I3450(ID47), 일렉트리파이드 GV70 95440-DS000·DS010(ID47), GV80 쿠페 95440-T6AA0(ID4A) 품번을 새로 확인.",
'',
'※ 조사하며 확인된 주요 사항',
"- 연식별 검색 기록은 research/hyundai/yearly.md, 예전 파일 내용은 research/hyundai/old_rows.tsv에 있습니다.",
"- 아반떼 XD는 2005년형(2004.7.1)부터, 싼타페 SM은 2005년형(2004.8.11)부터, 스타렉스는 뉴 스타렉스(2004.1.29)부터 이모빌라이저가 기본 적용됐습니다. 베르나 트랜스폼(2009.6)은 원가절감으로 이모빌라이저가 빠졌습니다.",
"- 아반떼 HD 리모컨은 2008년형까지 447MHz, 2009년형부터 433MHz입니다. 싼타페 CM 스마트키 95440-2B800과 베라크루즈 스마트키는 한국형 447MHz입니다.",
"- 그랜드 스타렉스·포터2는 이모빌라이저 사양과 비사양이 공존합니다(키·폴딩키 품번이 따로 있음).",
"- 칩 세대 흐름: ID46(Hitag2, ~2015년경) → ID8A(DST-AES, LF·AD·i30 PD·i40 후기) / ID47(Hitag3, IG·TL·TM·코나·제네시스) → ID4A(Hitag AES, GN7·DN8 디 엣지·MX5·아이오닉5·제네시스 부분변경) / ID6A(아반떼 CN7·캐스퍼·포터2 일렉트릭).",
]
for n in notes: g.append([n])
g.column_dimensions['A'].width = 150
for c in g['A']: c.font = Font(name='Arial', size=10); c.alignment = Alignment(wrap_text=False, vertical='top')
wb.save(OUT)

sys.path.insert(0, os.path.join(ROOT, 'research'))
from add_brand import add_brand; add_brand(OUT)                     # 브랜드 컬럼(이미 있으면 건너뜀)
from notation import apply_file; apply_file(OUT)                    # 표기 통일
sys.path.insert(0, os.path.join(ROOT, 'research', 'recheck')); import apply as recheck; recheck.apply_file(OUT)
print(len(rows), 'rows written')
