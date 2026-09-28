# 6차: 모델 세대별 심층 보완 — 빈 칸(키웨이·이모빌라이저·부품번호 등)을 세대 단위 규칙으로 채움
#   규칙(RULES): 파일·모델(정규식)·연식 범위·컬럼·값·출처. 해당 행의 그 칸이 비어 있을 때만 채우고, 출처 칸에 근거를 덧붙임.
#   결과는 overrides.json에 "r6": true 항목으로 저장(다시 실행하면 r6 항목만 새로 만듦) → 각 브랜드 build 때 apply.py가 자동 적용.
# 사용: python3 research/recheck/fill6.py        (overrides 갱신 + 엑셀에 바로 적용)
import os, re, sys, json, collections
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE); sys.path.insert(0, os.path.dirname(HERE))
from rules6 import RULES, FUNCS

CHIP = '칩코드 (예: ID46(PCF7936))'


def rows(path):
    ws = openpyxl.load_workbook(path).worksheets[0]
    hdr = [str(c.value).strip() if c.value else '' for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[hdr.index('모델명')]:
            yield dict(zip(hdr, r))


def build():
    out = collections.OrderedDict()
    for f in sorted({r['file'] for r in RULES} | {fn['file'] for fn in FUNCS}):
        path = os.path.join(ROOT, f)
        for row in rows(path):
            m, y = row['모델명'], int(str(row['연식'])[:4])
            done = set()
            cands = [r for r in RULES if r['file'] == f] + [dict(fn, dyn=True) for fn in FUNCS if fn['file'] == f]
            for r in cands:
                if not re.fullmatch(r['model'], m) or not (r.get('y0', 0) <= y <= r.get('y1', 9999)):
                    continue
                if r.get('chip') and not re.search(r['chip'], str(row.get(CHIP) or '')):
                    continue
                col = r['col']
                if col in done or (str(row.get(col) or '').strip() and not r.get('force')):
                    continue
                val = r['fn'](row) if r.get('dyn') else r['val']
                if not val:
                    continue
                done.add(col)
                key = (row['브랜드'], m, col, val, r['src'], r.get('flag'), r.get('note'))
                out.setdefault(key, []).append(y)
    ovs = []
    for (b, m, col, val, src, fl, note), ys in out.items():
        o = {'brand': b, 'model': m, 'years': sorted(set(ys)), 'set': {col: val}, 'append': {'출처': src}, 'r6': True}
        if fl is not None:
            o['flag'] = fl; o['note'] = note
        ovs.append(o)
    return ovs


if __name__ == '__main__':
    p = os.path.join(HERE, 'overrides.json')
    allo = json.load(open(p, encoding='utf-8'))
    cur = [o for o in allo if not o.get('r6') or o.get('guide')]
    old6 = [o for o in allo if o.get('r6') and not o.get('guide')]
    new = build()
    seen = {(o['brand'], o['model'], tuple(o['set'])) for o in new}
    # 이미 적용돼 빈 칸이 없어진 이전 r6 항목은 유지(다시 build해도 적용되도록)
    keep = [o for o in old6 if (o['brand'], o['model'], tuple(o['set'])) not in seen]
    json.dump(cur + keep + new, open(p, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    new = keep + new + [o for o in cur if o.get('guide') and o.get('r6')]
    print('r6 overrides:', len(new), 'cells:', sum(len(o['years']) for o in new))
    import apply
    for f in sorted({r['file'] for r in RULES} | {fn['file'] for fn in FUNCS}):
        print(f, apply.apply_file(os.path.join(ROOT, f), new))
