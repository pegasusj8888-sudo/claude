# 2005년 이후 주황·빨강 행 재조사 결과 적용 (research/recheck/overrides.json)
#   항목: {"brand","model","years":[...], "set":{컬럼:값}, "flag":"orange"|"red"|"", "note":"...", "delete":true}
#   - flag ""  → 색 지우고 비고 비움(확인됨)
#   - delete → 그 연식 행 삭제(국내 판매 없음 확인)
#   - 칩코드를 바꾸면 XT 호환 칸도 다시 계산
# 사용: python3 research/recheck/apply.py [파일…]   /   build_xlsx.py에서 apply_file(경로)
import os, sys, json, re
from copy import copy
HERE = os.path.dirname(os.path.abspath(__file__))
RESEARCH = os.path.dirname(HERE)
ROOT = os.path.dirname(RESEARCH)
sys.path.insert(0, RESEARCH)
from openpyxl.styles import PatternFill, Font

CHIP = '칩코드 (예: ID46(PCF7936))'
FILL = {'orange': 'FFFFD599', 'red': 'FFFFC7CE'}
FONT = {'orange': 'FF000000', 'red': 'FF9C0006'}
XTFILL = {'O': ('FFC6EFCE', 'FF006100'), '△': ('FFFFEB9C', 'FF9C6500'), 'X': ('FFFFC7CE', 'FF9C0006')}


def xt(chip):
    chip = chip or ''
    if not chip or (chip.startswith(('확인', '해당')) and 'ID' not in chip):
        return ('', '', '')
    new = any(k in chip for k in ('ID4A', 'ID47', 'ID49', 'ID8A', 'ID6A', 'ID75', 'ID88', 'AES'))
    old = any(k in chip for k in ('ID46', 'ID44', 'ID60', 'ID70', 'DST80', '4D70', '4D60x80', 'ID6E', 'ID4C'))
    if 'ID48' in chip:
        return ('△', '△', 'O') if old else ('△', '△', '△')
    if old and new: return ('△', '△', 'O')
    if old: return ('O', 'O', 'O')
    if new: return ('X', '△', 'O')
    return ('', '', '')


def load():
    return json.load(open(os.path.join(HERE, 'overrides.json'), encoding='utf-8'))


def apply_file(path, overrides=None):
    from add_brand import _load, _header_row
    ovs = overrides if overrides is not None else load()
    wb = _load(path)
    n = 0
    for ws in wb.worksheets:
        hr, _ = _header_row(ws)
        if hr is None:
            continue
        hdr = {str(c.value).strip(): c.column for c in ws[hr] if c.value}
        if CHIP not in hdr or '브랜드' not in hdr:
            continue
        dels = []
        for r in range(hr + 1, ws.max_row + 1):
            b = str(ws.cell(r, hdr['브랜드']).value or '')
            m = str(ws.cell(r, hdr['모델명']).value or '')
            y = str(ws.cell(r, hdr['연식']).value or '')
            for o in ovs:
                if o['brand'] != b or o['model'] != m or not any(y.startswith(str(yy)) for yy in o['years']):
                    continue
                if o.get('delete'):
                    dels.append(r); n += 1; break
                for col, val in o.get('set', {}).items():
                    ws.cell(r, hdr[col]).value = val if val != '' else None
                    n += 1
                for col in o.get('clear_fill', []):
                    ws.cell(r, hdr[col]).fill = PatternFill(fill_type=None)
                    n += 1
                if CHIP in o.get('set', {}) and 'XT27A/A66 호환' in hdr:
                    for c, v in zip(('XT27A/A66 호환', 'XT27B 호환', 'XT57B 호환'), xt(o['set'][CHIP])):
                        cell = ws.cell(r, hdr[c]); cell.value = v or None
                        if v in XTFILL:
                            cell.fill = PatternFill('solid', fgColor=XTFILL[v][0])
                            f = copy(cell.font); f.color = XTFILL[v][1]; cell.font = f
                        else:
                            cell.fill = PatternFill(fill_type=None)
                if 'flag' in o:
                    fl = o['flag']
                    for c in (CHIP, '이모빌라이저 시스템', '비고'):
                        if c not in hdr: continue
                        cell = ws.cell(r, hdr[c])
                        if c == '이모빌라이저 시스템' and fl and not str(cell.value or '').strip():
                            continue  # 빈 이모빌라이저 칸은 칠하지 않음
                        cell.fill = PatternFill('solid', fgColor=FILL[fl]) if fl else PatternFill(fill_type=None)
                        f = copy(cell.font); f.color = FONT.get(fl, 'FF000000'); cell.font = f
                    if '비고' in hdr:
                        ws.cell(r, hdr['비고']).value = (o.get('note') or None) if fl else None
                    n += 1
        guides = [o['guide'] for o in ovs if o.get('guide') and o['brand'] in
                  {str(ws.cell(r, hdr['브랜드']).value or '') for r in range(hr + 1, ws.max_row + 1)}]
        if guides and '안내' in wb.sheetnames:
            from openpyxl.styles import Alignment
            g = wb['안내']
            have = {c.value for c in g['A']}
            head = '※ 재조사로 확인한 근거 (색을 지운 행)'
            ref = g.cell(g.max_row, 1)
            for line in ([''] + [head] if head not in have else []) + ['- ' + t for t in guides if '- ' + t not in have]:
                r2 = g.max_row + 1
                g.cell(r2, 1).value = line or None
                g.cell(r2, 1)._style = copy(ref._style)
                g.cell(r2, 1).alignment = Alignment(wrap_text=False, vertical='top')
                n += 1
        # 빈 이모빌라이저 칸의 주황·빨강 채우기 지우기(불완전한 칸에만 색)
        if '이모빌라이저 시스템' in hdr:
            for r in range(hr + 1, ws.max_row + 1):
                cell = ws.cell(r, hdr['이모빌라이저 시스템'])
                if not str(cell.value or '').strip() and cell.fill.fill_type and \
                        str(cell.fill.fgColor.rgb).upper() in (FILL['orange'], FILL['red'], 'FFC7CE'):
                    cell.fill = PatternFill(fill_type=None); n += 1
        for r in reversed(dels):
            ws.delete_rows(r)
        if dels and ws.auto_filter.ref:
            ws.auto_filter.ref = re.sub(r'\d+$', str(ws.max_row), ws.auto_filter.ref)
    if n:
        wb.save(path)
    return n


if __name__ == '__main__':
    files = sys.argv[1:] or [os.path.join(ROOT, f) for f in sorted(os.listdir(ROOT)) if f.endswith('.xlsx')]
    ovs = load()
    for f in files:
        print(os.path.basename(f), apply_file(f, ovs))
