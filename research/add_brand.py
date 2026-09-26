# 모든 엑셀 파일의 데이터 시트 맨 앞에 '브랜드' 컬럼을 넣음(이미 있으면 건너뜀).
# 사용: python3 research/add_brand.py  (저장소 루트의 모든 .xlsx 처리)
#       build_xlsx.py에서는 add_brand(경로, 브랜드) 호출
import os, re, io, sys, zipfile
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill, Alignment
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.dimensions import ColumnDimension

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MODEL_ALIASES = ('모델명', '모델이름', '모델', '차종', '차명')
BRANDS = [  # 파일명에 들어 있는 글자 → 브랜드
    ('기아', '기아'), ('벤츠', '벤츠'), ('BMW', 'BMW'), ('아우디', '아우디'), ('audi', '아우디'),
    ('쉐보레', '쉐보레(GM대우)'), ('chevrolet', '쉐보레(GM대우)'), ('르노', '르노(르노삼성)'),
    ('KG모빌리티', 'KG모빌리티(쌍용)'), ('현대', '현대'), ('제네시스', '제네시스'),
]
NOTE = "- 맨 앞 '브랜드' 컬럼에 브랜드명을 적었습니다(검색 프로그램에서 '기아 K5'처럼 브랜드와 모델명을 함께 검색 가능)."


def brand_of(fname):
    for key, b in BRANDS:
        if key.lower() in fname.lower():
            return b
    return None


def _load(path):
    """한셀에서 저장한 파일(styles.xml 안 mc:AlternateContent)은 Fallback 서식으로 풀어서 읽음."""
    try:
        return openpyxl.load_workbook(path)
    except Exception:
        buf = io.BytesIO()
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info)
                if info.filename == 'xl/styles.xml':
                    s = data.decode('utf-8')
                    s = re.sub(r'<mc:AlternateContent\b[^>]*>.*?<mc:Fallback>(.*?)</mc:Fallback>\s*</mc:AlternateContent>',
                               r'\1', s, flags=re.S)
                    data = s.encode('utf-8')
                zout.writestr(info, data)
        buf.seek(0)
        return openpyxl.load_workbook(buf)


def _header_row(ws):
    for r in range(1, min(ws.max_row, 10) + 1):
        for c in range(1, min(ws.max_column, 30) + 1):
            if str(ws.cell(r, c).value or '').strip() in MODEL_ALIASES:
                return r, c
    return None, None


def _shift_ref(ref):
    """'B2' → 'C2', 'A1:O426' → 'A1:P426' (1열은 그대로 둬서 브랜드 칸 포함)"""
    def one(cell, keep_a):
        m = re.match(r'(\$?)([A-Z]+)(\$?\d*)$', cell)
        col = column_index_from_string(m.group(2))
        if keep_a and col == 1:
            return cell
        return m.group(1) + get_column_letter(col + 1) + m.group(3)
    parts = ref.split(':')
    return ':'.join(one(p, i == 0) for i, p in enumerate(parts))


def add_brand(path, brand=None):
    brand = brand or brand_of(os.path.basename(path))
    if not brand:
        print('브랜드 모름, 건너뜀:', path)
        return False
    wb = _load(path)
    changed = False
    for ws in wb.worksheets:
        hr, mc = _header_row(ws)
        if hr is None:
            continue
        if str(ws.cell(hr, 1).value or '').strip() == '브랜드':
            continue
        # 컬럼 너비 기억
        widths = {}
        for key, d in list(ws.column_dimensions.items()):
            if d.width:
                for i in range(d.min or column_index_from_string(key), (d.max or column_index_from_string(key)) + 1):
                    widths[i] = d.width
        freeze = ws.freeze_panes
        af = ws.auto_filter.ref
        ws.insert_cols(1)
        ws.column_dimensions.clear()
        for i, w in widths.items():
            L = get_column_letter(i + 1)
            ws.column_dimensions[L] = ColumnDimension(ws, index=L, width=w)
        ws.column_dimensions['A'] = ColumnDimension(ws, index='A', width=14)
        if freeze and freeze != 'A1':
            ws.freeze_panes = 'C' + re.sub(r'[A-Z]+', '', freeze) if freeze.startswith('B') else _shift_ref(freeze)
        if af:
            ws.auto_filter.ref = _shift_ref(af)
        model_col = mc + 1
        for r in range(hr, ws.max_row + 1):
            src = ws.cell(r, model_col)
            if r > hr and not str(src.value or '').strip():
                continue
            cell = ws.cell(r, 1)
            cell.value = '브랜드' if r == hr else brand
            cell._style = copy(src._style)
            if r > hr:
                cell.fill = PatternFill(fill_type=None)
            al = copy(src.alignment)
            cell.alignment = Alignment(horizontal=al.horizontal, vertical=al.vertical or 'center', wrap_text=False)
        changed = True
    # 모든 셀 줄바꿈 끄기(행 1줄 규칙)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.alignment.wrap_text:
                    al = copy(c.alignment)
                    al.wrap_text = False
                    c.alignment = al
                    changed = True
    if changed:
        # 안내 시트에 설명 한 줄 추가
        for ws in wb.worksheets:
            if ws.title != '안내':
                continue
            if any(NOTE == c.value for c in ws['A']):
                break
            pos = next((c.row for c in ws['A'] if str(c.value or '').strip() == '※ 구성'), None)
            at = pos + 1 if pos else ws.max_row + 1
            ref = ws.cell(at if at <= ws.max_row else ws.max_row, 1)
            ws.insert_rows(at)
            ws.cell(at, 1).value = NOTE
            ws.cell(at, 1)._style = copy(ref._style)
            ws.cell(at, 1).alignment = Alignment(wrap_text=False, vertical='top')
        wb.save(path)
    return changed


if __name__ == '__main__':
    files = sys.argv[1:] or [os.path.join(ROOT, f) for f in sorted(os.listdir(ROOT)) if f.endswith('.xlsx')]
    for f in files:
        print(os.path.basename(f), brand_of(os.path.basename(f)), '추가' if add_brand(f) else '변경 없음')
