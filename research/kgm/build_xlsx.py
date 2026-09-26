# research/kgm/rows.jsonl → KG모빌리티(쌍용)_models.xlsx (기아 트랜스폰더 DB와 동일 14컬럼)
import json, openpyxl, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
ROOT=os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
rows=[json.loads(l) for l in open(os.path.join(ROOT,'research','kgm','rows.jsonl'))]
order=[l.strip() for l in open(os.path.join(ROOT,'research','kgm','model_order.txt'))]
rows.sort(key=lambda r:(order.index(r['model']),int(r['year'][:4])))

def xt(chip):
    if chip.startswith('확인'): return ('','','')
    new=any(k in chip for k in ('ID4A','ID47','ID49','ID8A','AES'))
    old=any(k in chip for k in ('ID46','ID60','ID70','DST80','4D70'))
    if 'ID48' in chip:
        return ('△','△','O') if old else ('△','△','△')
    if old and new: return ('△','△','O')
    if old: return ('O','O','O')
    if new: return ('X','△','O')
    return ('','','')

HDR=['모델명','연식','XT27A/A66 호환','XT27B 호환','XT57B 호환','칩코드 (예: ID46(PCF7936))','키블레이드(부품번호)','키블레이드(키웨이)','카드키(부품번호)','스마트키(부품번호)','폴딩키(부품번호)','이모빌라이저 시스템','출처','비고']
wb=openpyxl.Workbook(); ws=wb.active; ws.title='KGM(쌍용) 트랜스폰더 DB'
thin=Side(style='thin',color='FFBFBFBF'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
ws.append(HDR)
for c in ws[1]:
    c.font=Font(name='Arial',size=11,bold=True,color='FFFFFFFF'); c.fill=PatternFill('solid',fgColor='FF7A1F1F')
    c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=bd
F=lambda c:PatternFill('solid',fgColor=c)
XT={'O':(F('FFC6EFCE'),'FF006100'),'△':(F('FFFFEB9C'),'FF9C6500'),'X':(F('FFFFC7CE'),'FF9C0006')}
FLAG={'orange':(F('FFFFD599'),'FF000000'),'red':(F('FFFFC7CE'),'FF9C0006'),'gen':(F('FFDDEBF7'),'FF000000')}
for r in rows:
    a,b,c=xt(r['chip'])
    ws.append([r['model'],r['year'],a,b,c,r['chip'],r['blade_pn'],r['keyway'],r['card'],r['smart'],r['fold'],r['immo'],r['src'],r['note']])
    i=ws.max_row
    for cell in ws[i]:
        cell.font=Font(name='Arial',size=10); cell.border=bd; cell.alignment=Alignment(vertical='center',wrap_text=True)
    for col in (3,4,5):
        cell=ws.cell(i,col); cell.alignment=Alignment(horizontal='center',vertical='center')
        if cell.value in XT: cell.fill,fc=XT[cell.value]; cell.font=Font(name='Arial',size=10,bold=True,color=fc)
    fl=r.get('flag')
    if fl:
        fill,fc=FLAG[fl]
        cols=(6,12,14) if fl in ('orange','red') else range(6,15)
        for col in cols:
            ws.cell(i,col).fill=fill; ws.cell(i,col).font=Font(name='Arial',size=10,color=fc)
for col,w in zip('ABCDEFGHIJKLMN',[30,14,12,11,11,26,26,22,12,36,12,22,48,48]): ws.column_dimensions[col].width=w
ws.row_dimensions[1].height=30; ws.freeze_panes='B2'; ws.auto_filter.ref=f'A1:N{ws.max_row}'

g=wb.create_sheet('안내')
notes=['이 표는 락스미스(자동차 키 제작/프로그래밍) 참고용입니다.','',
'※ 구성',
"- 컬럼은 '기아_트랜스폰더_칩코드_DBnew.xlsx'와 동일합니다. 모델마다 연식별로 한 행씩, 2000년식부터 기록했습니다.",
"- 칼리스타(1993-1997)·코란도 훼미리(1988-1996)는 2000년식 이전에 단종되어 행이 없습니다.",
"- 원본 목록 연식이 실제와 다른 모델은 실제 기준으로 적고 비고에 적었습니다(체어맨 H ~2014, 체어맨 W ~2017, 코란도 e-모션 2022~2023, 무쏘 EV 2025~, 렉스턴 스포츠는 2025.2 무쏘 스포츠로 차명 변경).",
"- '이모빌라이저 시스템': 구형은 VDO 이모빌라이저 박스(이모박스), 액티언·카이런·렉스턴 계열은 이모빌라이저 컨트롤 유닛(ICU, 87110-08B00/09000), 신형 스마트키 차량은 SKM(스마트키 모듈)입니다.",
"- 폴딩키(리모컨 일체형) 품번은 '폴딩키(부품번호)', 스마트키 트랜스미터 품번은 '스마트키(부품번호)' 칸에 적었습니다. 쌍용은 카드키가 없어 카드키 칸은 비었습니다.",
"- 'XT27A/A66·XT27B·XT57B 호환'은 기아 파일 규칙을 따랐습니다: ID60(4D60)·ID70/DST80(4D70) → O/O/O, ID48(Megamos Crypto) → △/△/△, ID48·ID60 병기 → △/△/O. 실제 적용 전 장비로 재확인하세요.",
"- 출처 칸은 해당 연식 검색에서 근거로 쓴 페이지입니다. 원문 페이지는 이 환경에서 열 수 없어 검색 결과 요약을 근거로 했습니다.",
'',
'※ 색상',
"- 빨간색: 확인 불가 — 토레스·토레스 EVX·액티언(J120)·무쏘 EV·이스타나는 키 칩 자료를 찾지 못했습니다.",
"- 주황색: 추정이거나 자료가 서로 다르거나 연식 범위 밖인 값 — 실물 키/VIN으로 재확인.",
'',
'※ 조사하며 확인된 주요 사항',
"- 무쏘·코란도·렉스턴(2002-2006): VDO 이모박스(MCU MC68HC05B16) + 메가모스 크립토(ID48) 트랜스폰더 — TMPro 모듈 85 자료. 코란도 2000-2006 블레이드 HYN10(JMA HY-5.P1).",
"- 액티언·카이런·렉스턴 계열: 4D-60 칩 키, 블레이드 SSY3(JMA SSA-2P). 순정 키 87170-32030·32020·08D50은 80bit, 87170-08B21은 40bit — 연식 구분 자료가 없어 키 품번으로 확인해야 합니다.",
"- 이모빌라이저 컨트롤 유닛(87110-08B00/09000) 적용: 렉스턴 2006.02~2017.04, 카이런 2005.05~2007.03, 액티언 2005.10~2007.03, 액티언 스포츠 2006.04~2007.03. 2007.4 이후 카이런·액티언 계열의 모듈 명칭은 확인하지 못했습니다.",
"- 순정 3버튼 폴딩키 87510-21100(코란도 C·로디우스 II/코란도 투리스모·티볼리 등): 80bit TI 칩(판매처 표기 ID70 DST80 / ID8E). 블레이드는 TOY48(remkeys)과 KI-7(3D Group, 블레이드 품번 7105121500)로 표기가 다릅니다.",
"- 코란도 C300 스마트키 DST80·HYN14R, 렉스턴 W 2013 스마트키 4D70 DST80(FCC DEO-MT FOBG02), G4 렉스턴 이모빌라이저 = SKM(87570-36010, 2017.07~)."]
for n in notes: g.append([n])
g.column_dimensions['A'].width=150
for c in g['A']: c.font=Font(name='Arial',size=10); c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save(os.path.join(ROOT,'KG모빌리티(쌍용)_models.xlsx'))
print(len(rows),'rows written')
