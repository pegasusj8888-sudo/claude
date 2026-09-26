# 기아 파일에 '키종류' 열 추가(칩코드와 키블레이드(부품번호) 사이) — 다른 브랜드 파일과 같은 구조로.
#   판정: 그 행에 실제로 적힌 부품번호·칩 조건을 근거로 막대키·폴딩키·스마트키·카드키를 모두 적음.
#     스마트키: 스마트키 부품번호, 칩 [스마트키], 비상키 '스마트키용'
#     카드키  : 카드키 부품번호, 칩 [카드키], 비상키 '카드키용'
#     폴딩키  : 폴딩키 부품번호, 칩 [폴딩키]/[플립키], 비상키 '폴딩키용'
#     막대키  : 칩·비상키에 이모빌라이저키/블랭킹키 표기. 부품 자료가 없는 구형 연식은 기본 막대키
#   MANUAL: 같은 세대 앞뒤 연식 자료로 보완한 행
# 사용: python3 research/kia/add_keytype.py   (이미 열이 있으면 값만 다시 계산)
import os, re, sys
from copy import copy
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(os.path.dirname(HERE))
KIA = os.path.join(ROOT, '기아_트랜스폰더_칩코드_DBnew.xlsx')
CHIP = '칩코드 (예: ID46(PCF7936))'; COL = '키종류'
TYPES = ['막대키', '폴딩키', '스마트키', '카드키']
MANUAL = {  # (모델명, 연식 앞 4자리 또는 None=전 연식) → 키종류
    ('프라이드 UB', None): '폴딩키,스마트키',          # UB 칩 ID6E-MA(폴딩키)·PCF7952(스마트키) 병존
    ('쏘울 (Soul, AM)', '2008'): '폴딩키,스마트키',     # 2008.9 출시, 2009년식과 같은 키
    ('카렌스 (Carens, UN)', '2013'): '폴딩키',          # UN 2007~2012와 같은 키
    ('K9 (K900/Quoris, KH)', '2018'): '스마트키,카드키',  # KH 2012~2017과 같은 키
}


def keytype(model, year, g):
    for (m, y), v in MANUAL.items():
        if m == model and (y is None or str(year)[:4] == y):
            return v
    chip, bl = g(CHIP), g('키블레이드(부품번호)')
    ts = set()
    if g('스마트키(부품번호)') or '스마트키' in chip or '스마트키용' in bl: ts.add('스마트키')
    if g('카드키(부품번호)') or '카드키' in chip or '카드키용' in bl: ts.add('카드키')
    if g('폴딩키(부품번호)') or re.search('폴딩|플립', chip) or '폴딩키용' in bl: ts.add('폴딩키')
    if re.search('이모빌라이저키|블랭킹|막대', chip + bl): ts.add('막대키')
    if not ts: ts.add('막대키')
    return ','.join(t for t in TYPES if t in ts)


def apply_file(path=KIA):
    wb = openpyxl.load_workbook(path); ws = wb.worksheets[0]
    hdr = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    if COL not in hdr:
        at = hdr[CHIP] + 1
        ws.insert_cols(at)
        # 열 너비·머리글 서식 옮기기
        from openpyxl.utils import get_column_letter as L
        for c in range(ws.max_column, at, -1):
            ws.column_dimensions[L(c)].width = ws.column_dimensions[L(c - 1)].width
        ws.column_dimensions[L(at)].width = 18
        for r in range(1, ws.max_row + 1):
            ws.cell(r, at)._style = copy(ws.cell(r, at + 1)._style)
            if r > 1: ws.cell(r, at).fill = openpyxl.styles.PatternFill(fill_type=None)
        ws.cell(1, at).value = COL
        if ws.auto_filter.ref:
            ws.auto_filter.ref = 'A1:' + L(ws.max_column) + str(ws.max_row)
        hdr = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    n = 0
    for r in range(2, ws.max_row + 1):
        m = ws.cell(r, hdr['모델명']).value
        if not m: continue
        g = lambda c: str(ws.cell(r, hdr[c]).value or '') if c in hdr else ''
        v = keytype(m, ws.cell(r, hdr['연식']).value, g)
        if ws.cell(r, hdr[COL]).value != v:
            ws.cell(r, hdr[COL]).value = v; n += 1
    wb.save(path)
    return n


if __name__ == '__main__':
    print('키종류', apply_file())
