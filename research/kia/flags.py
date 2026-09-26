# 기아 파일 불확실 행 색칠·비고 정리 (2026.9)
#  1) 불확실 행 → research/recheck/overrides.json 에 주황/빨강 항목 추가(재빌드해도 유지)
#  2) 비고 칸이 칠해지지 않은 행의 비고(출시·단종·부품번호 날짜 등 부가 설명) → '안내' 시트 '연식별 참고 메모'로 옮기고 비고 비움
# 사용: python3 research/kia/flags.py
import os, sys, json
from copy import copy
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, os.path.join(ROOT, 'research')); sys.path.insert(0, os.path.join(ROOT, 'research', 'recheck'))
from add_brand import _load, _header_row
import apply as recheck
from openpyxl.styles import Alignment

KIA = os.path.join(ROOT, '기아_트랜스폰더_칩코드_DBnew.xlsx')
CHIP = '칩코드 (예: ID46(PCF7936))'
R = lambda a, b: list(range(a, b + 1))
FOREIGN = '국내 부품번호·칩 원문 없음 — 해외 트랜스폰더 카탈로그의 칩 기준'
NEW = [
 # 자료 상충
 dict(brand='기아', model='K3 (Forte, BD)', years=[2019], set={CHIP: 'ID8A[자료 상충], ID47[자료 상충]'}, flag='orange',
      note='칩 자료 상충 — remotesandkeys는 2020 Forte TEXAS ID8A, dfwkeys4cars 종합표는 Forte 2019-2023 Hitag3 ID47'),
 dict(brand='기아', model='모닝 (Morning, TA/JA)', years=[2017], flag='orange',
      note='JA 칩 자료 상충 — ID6E-MA(DST80) 표기와 ID75(Texas AES) 표기가 갈림'),
 dict(brand='기아', model='모닝 (Morning, JA)', years=R(2018, 2025), flag='orange',
      note='JA 칩 자료 상충 — ID6E-MA(DST80) 표기와 ID75(Texas AES) 표기가 갈림'),
 dict(brand='기아', model='스포티지 (Sportage, SL)', years=[2013, 2014], flag='orange',
      note='플립키 ID6E-MA는 일부 자료에만 나옴(국제 판매처는 플립키 95430-3W200 ID46 PCF7936)'),
 # 해외 카탈로그 기준(국내 원문 없음)
 dict(brand='기아', model='스포티지 (Sportage, NB-Ⅶ)', years=R(1993, 1995), set={CHIP: '해당 없음[추정]'}, flag='orange',
      note='초기형 트랜스폰더 미탑재 추정 — 해외 카탈로그는 1996년부터 ID13'),
 dict(brand='기아', model='스포티지 (Sportage, NB-Ⅶ)', years=R(1996, 2002), flag='orange', note=FOREIGN + '(유럽 기준)'),
 dict(brand='기아', model='카니발 (Carnival, KV-II/GQ)', years=R(1998, 2004), flag='orange', note=FOREIGN + '(유럽 기준)'),
 dict(brand='기아', model='프라이드 JB', years=R(2006, 2010), flag='orange', note=FOREIGN),
 dict(brand='기아', model='프라이드 UB', years=[2012], flag='orange', note=FOREIGN),
 dict(brand='기아', model='모닝 (Morning, SA)', years=R(2005, 2009), flag='orange', note=FOREIGN),
 dict(brand='기아', model='스포티지 (Sportage, JE/KM)', years=R(2005, 2009), flag='orange', note=FOREIGN),
 dict(brand='기아', model='쏘렌토 (Sorento, BL)', years=R(2003, 2008), flag='orange', note=FOREIGN),
 dict(brand='기아', model='카렌스 (Carens, RS)', years=[2000, 2001, 2003, 2004, 2005], flag='orange', note=FOREIGN),
 dict(brand='기아', model='카렌스 (Carens, UN)', years=R(2008, 2011), flag='orange', note=FOREIGN),
 # 봉고3: 칩·이모빌라이저 칸만 칠해져 있고 비고에 이유가 없던 행
 dict(brand='기아', model='봉고3 (Bongo3, J)', years=R(2003, 2024), flag='orange',
      note='엔진·ECU 사양별 4D(81996-4E010)·46(81996-4E020) 키 구분 연식 매칭 원문 없음'),
 # 신형(부분변경·후속) 자료 없음
 dict(brand='기아', model='스토닉', years=[2025, 2026], flag='orange',
      note='2차 페이스리프트(2025.9) 이후 스마트키 개정 여부 원문 없음 — 기존 95440-H8000(ID8A) 기준 추정'),
 dict(brand='기아', model='셀토스 (SP2 PE)', years=[2026], flag='orange',
      note='디 올 뉴 셀토스(SP3, 2026.1) 스마트키 품번·칩 원문 없음 — SP2 PE 기준'),
 dict(brand='기아', model='니로 (Niro, SG2/SG2 PE)', years=[2026], flag='orange',
      note='더 뉴 니로(SG2 PE, 2026.3) 스마트키 품번·칩 원문 없음 — 초기형 기준'),
 # 빈 칩 채우기(같은 품번 근거 있음)
 dict(brand='기아', model='옵티마 (Optima, MS)', years=[2000], set={CHIP: 'ID60(4D60)'}),
 dict(brand='기아', model='오피러스 (Opirus, GH)', years=[2003], set={CHIP: 'ID60(4D60)', '키블레이드(부품번호)': '81996-3FA10[막대키]'}, flag=''),
 dict(brand='기아', model='쏘울 (Soul, AM)', years=[2008], set={CHIP: 'ID46(PCF7952A)'}),
 # 확인 불가
 dict(brand='기아', model='봉고 (1세대)', years=[2000, 2002], set={CHIP: '확인 불가'}, flag='red',
      note='칩 종류를 확인할 국내외 자료 없음'),
] + [dict(brand='기아', model=m, years=[y], set={CHIP: '확인 불가'}, flag='red',
          note='연식을 나누지 않은 예전 목록 행 — 칩 조사 대상에서 빠져 있음')
     for m, y in [('크레도스 (Credos/Credo)', '1995-1997'), ('크레도스 (Credos/Credo)', '1998-2001'),
                  ('세피아 II (Sephia)', '1997-2000'), ('스펙트라 (Spectra)', '2000-2003'), ('스펙트라 (Spectra, LD)', '2004-2009'),
                  ('쎄라토 (Cerato, LD)', '2004-2008'), ('아벨라 (Avella)', '1994-1999'), ('포텐샤 (Potentia)', '1992-2002'),
                  ('엔터프라이즈 (Enterprise)', '1997-2002'), ('비스토 (Visto)', '1997-2004')]] + [
 dict(brand='BMW', model='3시리즈 (E90/E91/E92/E93, 5세대)', years=[2010], flag='orange',
      note='칩 부품번호 PCF7943·PCF7944는 일부 자료에만 나옴(주 자료는 PCF7945)'),
 dict(brand='BMW', model='X4 (F26, 1세대)', years=[2016], flag='orange',
      note='칩 부품번호 PCF7945P는 일부 자료에만 나옴(주 자료는 PCF7953)'),
]

def move_notes(path):
    """비고 칸이 칠해져 있지 않은데 글이 있는 행 → 글을 '안내' 시트 '연식별 참고 메모'로 옮기고 비고 비움."""
    wb = _load(path); ws = wb.worksheets[0]
    hr, _ = _header_row(ws)
    hdr = {str(c.value).strip(): c.column for c in ws[hr] if c.value}
    COL = {'FFFFD599', 'FFFFC7CE', 'FFC7CE'}
    moved = []
    for r in range(hr + 1, ws.max_row + 1):
        note = ws.cell(r, hdr['비고'])
        if not str(note.value or '').strip():
            continue
        if note.fill.fill_type and str(note.fill.fgColor.rgb).upper() in COL:
            continue
        moved.append('- %s %s: %s' % (ws.cell(r, hdr['모델명']).value, ws.cell(r, hdr['연식']).value, str(note.value).strip(' —')))
        note.value = None
    g = wb['안내']
    head = '※ 연식별 참고 메모 (예전 비고 칸 내용 — 색 없는 행이라 비고에서 옮김)'
    have = {c.value for c in g['A']}
    ref = g.cell(g.max_row, 1)
    for line in ([''] + [head] if head not in have else []) + [m for m in moved if m not in have]:
        r2 = g.max_row + 1
        g.cell(r2, 1).value = line or None
        g.cell(r2, 1)._style = copy(ref._style)
        g.cell(r2, 1).alignment = Alignment(wrap_text=False, vertical='top')
    wb.save(path)
    return len(moved)


if __name__ == '__main__':
    p = os.path.join(ROOT, 'research', 'recheck', 'overrides.json')
    ovs = json.load(open(p, encoding='utf-8'))
    key = lambda o: (o['brand'], o['model'], tuple(map(str, o['years'])))
    have = {key(o) for o in ovs}
    ovs += [o for o in NEW if key(o) not in have]
    json.dump(ovs, open(p, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('overrides', len(ovs))
    print('비고 → 안내', move_notes(KIA))
    print('기아', recheck.apply_file(KIA, ovs))
