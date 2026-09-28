# 수입 5개 브랜드(렉서스·폭스바겐·볼보·랜드로버·포드) 공통 생성기
#   research/import5/<brand>.py 의 BRAND·FILE·TITLE·MODELS·GUIDE 를 읽어
#   연식별 검색 기록(yearly_<brand>.md)·rows_<brand>.jsonl·엑셀을 만든다.
# 사용: python3 research/import5/build.py lexus
import json, os, sys, importlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE); sys.path.insert(0, os.path.join(ROOT, 'research'))
from finalize import process

KEYS = ('chip', 'immo', 'keyway', 'ktype_hint', 'smart', 'fold', 'card', 'blade_pn', 'src', 'flag', 'note')


def expand(cfg):
    """MODELS: [(모델명, 목록연식, [세그먼트...])]
       세그먼트 = dict(y=(시작,끝), chip=…, … , split={연식:[{…},{…}]}) — split은 한 연식을 2행으로(생산 시기 구분)"""
    rows, log = [], []
    for m, listed, segs in cfg.MODELS:
        years = sorted({y for s in segs for y in range(s['y'][0], s['y'][1] + 1)})
        last = years[-1]
        ended = cfg.ENDED.get(m, True)
        for y in years:
            s = [s for s in segs if s['y'][0] <= y <= s['y'][1]][-1]
            parts = s.get('split', {}).get(y) or [{}]
            for sub in parts:
                d = {k: s.get(k, '') for k in KEYS}
                d.update(sub)
                d['model'] = m
                d['year'] = str(y) + ('(단종)' if ended and y == last and last < 2026 else '')
                if not d['flag']:
                    d.pop('flag'); d['note'] = ''
                rows.append(d)
                res = '; '.join(f'{k} {d[k]}' for k in ('chip', 'ktype_hint', 'smart', 'fold', 'card', 'blade_pn', 'immo') if d.get(k))
                log.append('|'.join((m, str(y), res.replace('|', '/'), d['src'].replace('|', '/'))))
    for r in rows:
        r.pop('logres', None); r.pop('res', None); r.pop('y', None); r.pop('split', None)
    return rows, log


def xt(chip):
    if chip.startswith(('확인', '해당')) or not chip: return ('', '', '')
    new = any(k in chip for k in ('ID4A', 'ID47', 'ID49', 'ID8A', 'ID88', 'AES', 'ID6A'))
    old = any(k in chip for k in ('ID46', 'ID44', 'ID60', 'ID63', 'ID70', 'ID4C', 'ID67', 'ID68', 'ID72', 'ID33', 'ID73', 'ID13', '4D60x80'))
    if 'ID48' in chip:
        return ('△', '△', 'O') if old else ('△', '△', '△')
    if old and new: return ('△', '△', 'O')
    if old: return ('O', 'O', 'O')
    if new: return ('X', '△', 'O')
    return ('', '', '')


HDR = ['모델명', '연식', 'XT27A/A66 호환', 'XT27B 호환', 'XT57B 호환', '칩코드 (예: ID46(PCF7936))', '키종류', '키블레이드(부품번호)', '키블레이드(키웨이)', '카드키(부품번호)', '스마트키(부품번호)', '폴딩키(부품번호)', '이모빌라이저 시스템', '출처', '비고']


def write_xlsx(cfg, rows):
    order = [m for m, _, _ in cfg.MODELS]
    rows.sort(key=lambda r: (order.index(r['model']), int(r['year'][:4])))
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = cfg.TITLE
    thin = Side(style='thin', color='FFBFBFBF'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(HDR)
    for c in ws[1]:
        c.font = Font(name='Arial', size=11, bold=True, color='FFFFFFFF'); c.fill = PatternFill('solid', fgColor='FF7A1F1F')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False); c.border = bd
    F = lambda c: PatternFill('solid', fgColor=c)
    XT = {'O': (F('FFC6EFCE'), 'FF006100'), '△': (F('FFFFEB9C'), 'FF9C6500'), 'X': (F('FFFFC7CE'), 'FF9C0006')}
    FLAG = {'orange': (F('FFFFD599'), 'FF000000'), 'red': (F('FFFFC7CE'), 'FF9C0006')}
    for r in rows:
        a, b, c = xt(r['chip'])
        ws.append([r['model'], r['year'], a, b, c, r['chip'], r['ktype'], r['blade_pn'], r['keyway'], r['card'], r['smart'], r['fold'], r['immo'], r['src'], r['note']])
        i = ws.max_row
        for cell in ws[i]:
            cell.font = Font(name='Arial', size=10); cell.border = bd; cell.alignment = Alignment(vertical='center', wrap_text=False)
        for col in (3, 4, 5):
            cell = ws.cell(i, col); cell.alignment = Alignment(horizontal='center', vertical='center')
            if cell.value in XT: cell.fill, fc = XT[cell.value]; cell.font = Font(name='Arial', size=10, bold=True, color=fc)
        fl = r.get('flag')
        if fl:
            fill, fc = FLAG[fl]
            for col in (6, 13, 15):
                if col == 13 and not r['immo']: continue
                ws.cell(i, col).fill = fill; ws.cell(i, col).font = Font(name='Arial', size=10, color=fc)
    for col, w in zip('ABCDEFGHIJKLMNO', [30, 14, 12, 11, 11, 30, 16, 30, 22, 22, 40, 26, 30, 60, 48]): ws.column_dimensions[col].width = w
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = f'A1:O{ws.max_row}'
    g = wb.create_sheet('안내')
    for n in cfg.GUIDE: g.append([n])
    g.column_dimensions['A'].width = 150
    for c in g['A']: c.font = Font(name='Arial', size=10); c.alignment = Alignment(wrap_text=False, vertical='top')
    path = os.path.join(ROOT, cfg.FILE)
    wb.save(path)
    from add_brand import add_brand; add_brand(path)
    from notation import apply_file; apply_file(path)
    return path


if __name__ == '__main__':
    b = sys.argv[1]
    cfg = importlib.import_module(b)
    rows, log = expand(cfg)
    open(os.path.join(HERE, f'yearly_{b}.md'), 'w', encoding='utf-8').write(
        f'# {cfg.BRAND} 연식별 검색 기록 (모델|연식|결과|출처)\n' + '\n'.join(log) + '\n')
    rows = process(b, rows)
    open(os.path.join(HERE, f'rows_{b}.jsonl'), 'w', encoding='utf-8').write(''.join(json.dumps(r, ensure_ascii=False) + '\n' for r in rows))
    p = write_xlsx(cfg, rows)
    from collections import Counter
    print(p, len(rows), Counter(r.get('flag') for r in rows))
